"""Cadastro em lote da Organização a partir de planilhas do Excel (.xlsx).

As planilhas seguem as mesmas colunas dos arquivos exportados em JSON
(alunosExportados.json, professoresExportados.json, turmasExportadas.json,
disciplinasExportadas.json):
a primeira linha traz os nomes das colunas e cada linha seguinte é um registro.
"""
import io
import re
import zipfile

from django.core.management.color import no_style
from django.db import connection, transaction
from django.http import HttpResponse, JsonResponse
from django.views.decorators.csrf import csrf_exempt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException

from .models import AtravessaPor, Disciplina, Estudante, Professor
from .views import _coordenador_logado

COLUNAS = {
    "alunos": [
        "id", "posicao_ordem", "nome_completo", "escola", "modo_de_ensino",
        "serie", "periodo", "curso", "turma", "senha", "ip",
    ],
    "professores": ["id", "nome_completo", "senha", "ip"],
    "turmas": ["id", "professor", "escola", "turma", "etapa", "disciplina_lecionada"],
    "disciplinas": ["id", "nome_disciplina", "professores"],
}

COLUNAS_OBRIGATORIAS = {
    "alunos": ["nome_completo"],
    "professores": ["nome_completo"],
    "turmas": ["professor", "turma", "disciplina_lecionada"],
    "disciplinas": ["nome_disciplina"],
}

TAMANHO_MAXIMO = 4 * 1024 * 1024  # limite de corpo de requisição da Vercel é 4,5 MB
MAXIMO_LINHAS = 5000
MAXIMO_ERROS_RETORNADOS = 100
SENHA_PADRAO_ALUNO = "12345678"


class ErroPlanilha(Exception):
    pass


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _inteiro(valor, nome_coluna):
    texto = _texto(valor)
    if not texto:
        return None
    try:
        return int(texto)
    except ValueError:
        raise ValueError(f"a coluna '{nome_coluna}' deve ser um número inteiro (valor: '{texto}').")


def _ler_planilha(arquivo, tipo):
    """Retorna (linhas, colunas_presentes); cada linha é (número da linha, dict)."""
    if not arquivo.name.lower().endswith(".xlsx"):
        raise ErroPlanilha("Envie um arquivo do Excel no formato .xlsx.")
    if arquivo.size > TAMANHO_MAXIMO:
        raise ErroPlanilha("Arquivo muito grande. O limite é de 4 MB.")

    try:
        workbook = load_workbook(io.BytesIO(arquivo.read()), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise ErroPlanilha("Não foi possível ler o arquivo. Confirme que é uma planilha .xlsx válida.")

    try:
        linhas_brutas = workbook.worksheets[0].iter_rows(values_only=True)
        cabecalho = next(linhas_brutas, None)
        if not cabecalho:
            raise ErroPlanilha("A planilha está vazia.")

        indice_por_coluna = {}
        for indice, celula in enumerate(cabecalho):
            nome = _texto(celula).lower().replace(" ", "_")
            if nome in COLUNAS[tipo] and nome not in indice_por_coluna:
                indice_por_coluna[nome] = indice

        faltando = [c for c in COLUNAS_OBRIGATORIAS[tipo] if c not in indice_por_coluna]
        if faltando:
            raise ErroPlanilha(
                f"Coluna(s) obrigatória(s) ausente(s) na primeira linha: {', '.join(faltando)}. "
                f"Colunas esperadas: {', '.join(COLUNAS[tipo])}."
            )

        linhas = []
        for numero, valores in enumerate(linhas_brutas, start=2):
            registro = {
                coluna: (valores[indice] if indice < len(valores) else None)
                for coluna, indice in indice_por_coluna.items()
            }
            if not any(_texto(v) for v in registro.values()):
                continue
            linhas.append((numero, registro))
            if len(linhas) > MAXIMO_LINHAS:
                raise ErroPlanilha(f"A planilha passa do limite de {MAXIMO_LINHAS} linhas por envio.")
    finally:
        workbook.close()

    if not linhas:
        raise ErroPlanilha("A planilha não tem nenhuma linha de dados abaixo do cabeçalho.")

    return linhas, set(indice_por_coluna)


def _reiniciar_sequencia(modelo):
    # Registros criados com id explícito não avançam a sequência do Postgres;
    # sem isso, o próximo cadastro normal tentaria reutilizar um id já ocupado.
    with connection.cursor() as cursor:
        for sql in connection.ops.sequence_reset_sql(no_style(), [modelo]):
            cursor.execute(sql)


def _ler_id(registro, numero, ids_vistos, erros):
    """Retorna (ok, id). Registra o erro da linha quando o id é inválido ou repetido."""
    try:
        registro_id = _inteiro(registro.get("id"), "id")
    except ValueError as erro:
        erros.append(f"Linha {numero}: {erro}")
        return False, None

    if registro_id is not None:
        if registro_id in ids_vistos:
            erros.append(f"Linha {numero}: o id {registro_id} aparece mais de uma vez na planilha.")
            return False, None
        ids_vistos.add(registro_id)

    return True, registro_id


def _importar_alunos(linhas, colunas, coordenador):
    campos_texto = [c for c in ("escola", "modo_de_ensino", "serie", "periodo", "curso", "turma") if c in colunas]
    existentes = Estudante.objects.in_bulk()
    nomes = {(a.nome_completo or "").upper() for a in existentes.values()}

    novos, atualizados, erros, ids_vistos = [], [], [], set()
    ignorados = 0

    for numero, registro in linhas:
        ok, aluno_id = _ler_id(registro, numero, ids_vistos, erros)
        if not ok:
            continue

        nome = _texto(registro.get("nome_completo")).upper()
        if not nome:
            erros.append(f"Linha {numero}: a coluna 'nome_completo' é obrigatória.")
            continue

        try:
            posicao_ordem = _inteiro(registro.get("posicao_ordem"), "posicao_ordem")
        except ValueError as erro:
            erros.append(f"Linha {numero}: {erro}")
            continue

        senha = _texto(registro.get("senha"))
        ip = _texto(registro.get("ip"))

        if aluno_id in existentes:
            aluno = existentes[aluno_id]
            aluno.nome_completo = nome
            for campo in campos_texto:
                setattr(aluno, campo, _texto(registro.get(campo)))
            if "posicao_ordem" in colunas:
                aluno.posicao_ordem = posicao_ordem
            if senha:
                aluno.senha = senha
            if ip:
                aluno.ip = ip
            atualizados.append(aluno)
            continue

        if nome in nomes:
            ignorados += 1
            continue
        nomes.add(nome)

        novos.append(Estudante(
            id=aluno_id,
            posicao_ordem=posicao_ordem,
            nome_completo=nome,
            escola=_texto(registro.get("escola")) or (coordenador.escola or ""),
            modo_de_ensino=_texto(registro.get("modo_de_ensino")),
            serie=_texto(registro.get("serie")),
            periodo=_texto(registro.get("periodo")),
            curso=_texto(registro.get("curso")),
            turma=_texto(registro.get("turma")),
            senha=senha or SENHA_PADRAO_ALUNO,
            ip=ip or None,
        ))

    with transaction.atomic():
        Estudante.objects.bulk_create(novos)
        if atualizados:
            Estudante.objects.bulk_update(
                atualizados,
                ["nome_completo", "posicao_ordem", "senha", "ip", *campos_texto],
                batch_size=200,
            )
        _reiniciar_sequencia(Estudante)

    return len(novos), len(atualizados), ignorados, erros


def _importar_professores(linhas, colunas, coordenador):
    existentes = Professor.objects.in_bulk()
    id_por_nome = {(p.nome_completo or "").upper(): p.id for p in existentes.values()}

    novos, atualizados, erros, ids_vistos = [], [], [], set()
    nomes_novos = set()
    ignorados = 0

    for numero, registro in linhas:
        ok, professor_id = _ler_id(registro, numero, ids_vistos, erros)
        if not ok:
            continue

        nome = _texto(registro.get("nome_completo")).upper()
        if not nome:
            erros.append(f"Linha {numero}: a coluna 'nome_completo' é obrigatória.")
            continue

        senha = _texto(registro.get("senha"))
        ip = _texto(registro.get("ip"))

        if professor_id in existentes:
            dono_do_nome = id_por_nome.get(nome)
            if nome in nomes_novos or (dono_do_nome is not None and dono_do_nome != professor_id):
                erros.append(f"Linha {numero}: já existe outro professor com o nome '{nome}'.")
                continue

            professor = existentes[professor_id]
            id_por_nome.pop((professor.nome_completo or "").upper(), None)
            id_por_nome[nome] = professor_id
            professor.nome_completo = nome
            if senha:
                professor.senha = senha
            if ip:
                professor.ip = ip
            atualizados.append(professor)
            continue

        if nome in id_por_nome or nome in nomes_novos:
            ignorados += 1
            continue
        nomes_novos.add(nome)

        novos.append(Professor(
            id=professor_id,
            nome_completo=nome,
            senha=senha or f"{nome.split()[0]}123",
            ip=ip or None,
        ))

    with transaction.atomic():
        Professor.objects.bulk_create(novos)
        if atualizados:
            Professor.objects.bulk_update(atualizados, ["nome_completo", "senha", "ip"], batch_size=200)
        _reiniciar_sequencia(Professor)

    return len(novos), len(atualizados), ignorados, erros


def _importar_turmas(linhas, colunas, coordenador):
    professores = Professor.objects.in_bulk()
    id_por_nome = {(p.nome_completo or "").upper(): p.id for p in professores.values()}
    existentes = AtravessaPor.objects.in_bulk()
    chaves = {(v.professor_id, v.turma, v.disciplina_lecionada) for v in existentes.values()}

    novos, atualizados, erros, ids_vistos = [], [], [], set()
    ignorados = 0

    for numero, registro in linhas:
        ok, vinculo_id = _ler_id(registro, numero, ids_vistos, erros)
        if not ok:
            continue

        # A coluna 'professor' aceita o id (como no JSON exportado) ou o nome completo.
        referencia = _texto(registro.get("professor"))
        turma = _texto(registro.get("turma"))
        disciplina = _texto(registro.get("disciplina_lecionada"))

        if not referencia or not turma or not disciplina:
            erros.append(
                f"Linha {numero}: as colunas 'professor', 'turma' e 'disciplina_lecionada' são obrigatórias."
            )
            continue

        professor_id = int(referencia) if referencia.isdigit() else id_por_nome.get(referencia.upper())
        if professor_id not in professores:
            erros.append(
                f"Linha {numero}: professor '{referencia}' não está cadastrado. "
                f"Cadastre o professor antes de importar a turma."
            )
            continue

        escola = _texto(registro.get("escola"))
        etapa = _texto(registro.get("etapa"))

        if vinculo_id in existentes:
            vinculo = existentes[vinculo_id]
            vinculo.professor_id = professor_id
            vinculo.turma = turma
            vinculo.disciplina_lecionada = disciplina
            if "escola" in colunas:
                vinculo.escola = escola
            if "etapa" in colunas:
                vinculo.etapa = etapa
            atualizados.append(vinculo)
            continue

        chave = (professor_id, turma, disciplina)
        if chave in chaves:
            ignorados += 1
            continue
        chaves.add(chave)

        novos.append(AtravessaPor(
            id=vinculo_id,
            professor_id=professor_id,
            escola=escola or (coordenador.escola or ""),
            turma=turma,
            etapa=etapa,
            disciplina_lecionada=disciplina,
        ))

    with transaction.atomic():
        AtravessaPor.objects.bulk_create(novos)
        if atualizados:
            AtravessaPor.objects.bulk_update(
                atualizados,
                ["professor", "escola", "turma", "etapa", "disciplina_lecionada"],
                batch_size=200,
            )
        _reiniciar_sequencia(AtravessaPor)

    return len(novos), len(atualizados), ignorados, erros


def _importar_disciplinas(linhas, colunas, coordenador):
    professores = Professor.objects.in_bulk()
    id_por_nome_professor = {(p.nome_completo or "").upper(): p.id for p in professores.values()}
    existentes = Disciplina.objects.in_bulk()
    nomes = {(d.nome_disciplina or "").upper() for d in existentes.values()}

    # Cada item é (disciplina, ids dos professores ou None quando a coluna não veio).
    novos, atualizados, erros, ids_vistos = [], [], [], set()
    ignorados = 0

    for numero, registro in linhas:
        ok, disciplina_id = _ler_id(registro, numero, ids_vistos, erros)
        if not ok:
            continue

        nome = _texto(registro.get("nome_disciplina")).upper()
        if not nome:
            erros.append(f"Linha {numero}: a coluna 'nome_disciplina' é obrigatória.")
            continue

        # A coluna 'professores' é uma lista separada por vírgula ou ponto e
        # vírgula, com o id (como no JSON exportado) ou o nome completo de cada um.
        ids_professores = None
        if "professores" in colunas:
            ids_professores = []
            desconhecidos = []
            for referencia in re.split(r"[,;]", _texto(registro.get("professores"))):
                referencia = referencia.strip()
                if not referencia:
                    continue
                professor_id = (
                    int(referencia) if referencia.isdigit() else id_por_nome_professor.get(referencia.upper())
                )
                if professor_id in professores:
                    ids_professores.append(professor_id)
                else:
                    desconhecidos.append(referencia)
            if desconhecidos:
                erros.append(
                    f"Linha {numero}: professor(es) não cadastrado(s): {', '.join(desconhecidos)}."
                )
                continue

        if disciplina_id in existentes:
            disciplina = existentes[disciplina_id]
            disciplina.nome_disciplina = nome
            atualizados.append((disciplina, ids_professores))
            continue

        if nome in nomes:
            ignorados += 1
            continue
        nomes.add(nome)

        novos.append((Disciplina(id=disciplina_id, nome_disciplina=nome), ids_professores))

    with transaction.atomic():
        for disciplina, ids_professores in novos + atualizados:
            disciplina.save()
            if ids_professores is not None:
                disciplina.professores.set(ids_professores)
        _reiniciar_sequencia(Disciplina)

    return len(novos), len(atualizados), ignorados, erros


IMPORTADORES = {
    "alunos": _importar_alunos,
    "professores": _importar_professores,
    "turmas": _importar_turmas,
    "disciplinas": _importar_disciplinas,
}


@csrf_exempt
def importar_planilha(request, tipo):
    if request.method != "POST":
        return JsonResponse({"message": "Método não permitido."}, status=405)

    coordenador = _coordenador_logado(request)
    if not coordenador:
        return JsonResponse({"message": "Não autenticado."}, status=401)

    if tipo not in IMPORTADORES:
        return JsonResponse({"message": "Tipo de importação desconhecido."}, status=404)

    arquivo = request.FILES.get("arquivo")
    if not arquivo:
        return JsonResponse({"message": "Selecione uma planilha para enviar."}, status=400)

    try:
        linhas, colunas = _ler_planilha(arquivo, tipo)
    except ErroPlanilha as erro:
        return JsonResponse({"message": str(erro)}, status=400)

    criados, atualizados, ignorados, erros = IMPORTADORES[tipo](linhas, colunas, coordenador)

    return JsonResponse({
        "message": (
            f"{criados} registro(s) criado(s), {atualizados} atualizado(s), "
            f"{ignorados} ignorado(s) por já existirem e {len(erros)} linha(s) com erro."
        ),
        "criados": criados,
        "atualizados": atualizados,
        "ignorados": ignorados,
        "total_erros": len(erros),
        "erros": erros[:MAXIMO_ERROS_RETORNADOS],
    })


@csrf_exempt
def modelo_planilha(request, tipo):
    if not _coordenador_logado(request):
        return JsonResponse({"message": "Não autenticado."}, status=401)

    if tipo not in COLUNAS:
        return JsonResponse({"message": "Tipo de importação desconhecido."}, status=404)

    workbook = Workbook()
    planilha = workbook.active
    planilha.title = tipo
    planilha.append(COLUNAS[tipo])
    for celula in planilha[1]:
        celula.font = Font(bold=True)
    for coluna in planilha.columns:
        planilha.column_dimensions[coluna[0].column_letter].width = 24

    conteudo = io.BytesIO()
    workbook.save(conteudo)

    response = HttpResponse(
        conteudo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="modelo_{tipo}.xlsx"'
    return response
